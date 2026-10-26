"""Expense claims tests. Run: python3 -m unittest discover -s secretary"""

import json
import os
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import claims as C  # noqa: E402

TODAY = date.today()


def d(n: int) -> str:
    """The date n days from today."""
    return (TODAY + timedelta(days=n)).isoformat()


D1, D2 = d(-22), d(-21)  # two receipt days on a trip


def args(**kw):
    base = dict(date=D1, merchant="Grab", amount=182000.0, currency="VND", category="transport", description="Airport to hotel",
                trip=None, no_trip=False, home=None, receipt=None, note=None, force=False)
    base.update(kw)
    return SimpleNamespace(**base)


class Claims(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.saved = (C.DATA, C.HOME, C.drop_receipt, C.get_receipt, C.sync.config, C.sync.token, C.sync.post)
        C.DATA = Path(self.tmp.name)
        C.HOME = "SGD"  # the tests' home currency, whatever settings.json says
        # cached ECB rates, so no network: home currency per unit
        (C.DATA / "fx_cache.json").write_text(json.dumps({f"VND:SGD:{D1}": {"rate": 0.0000496, "date": D1, "source": "test"},
                                                          f"VND:SGD:{D2}": {"rate": 0.0000495, "date": D2, "source": "test"},
                                                          f"JPY:SGD:{D2}": {"rate": 0.00862, "date": D2, "source": "test"},
                                                          f"VND:EUR:{D1}": {"rate": 0.0000346, "date": D1, "source": "test"}}))

    def tearDown(self):
        C.DATA, C.HOME, C.drop_receipt, C.get_receipt, C.sync.config, C.sync.token, C.sync.post = self.saved
        self.tmp.cleanup()

    def test_add_converts_tags_the_open_trip_and_catches_duplicates(self):
        C.save_state({"trip": "Hanoi trip"})
        r = C.add(args())
        self.assertEqual(r["home"], "9.03")
        self.assertEqual(r["trip"], "Hanoi trip")
        self.assertEqual(r["status"], "to_claim")
        with self.assertRaises(SystemExit):
            C.add(args())  # same day, merchant, amount
        self.assertEqual(len(C.load()), 1)
        C.add(args(force=True))
        self.assertEqual(len(C.load()), 2)
        # the card's own charge wins over the ECB rate; --no-trip leaves the trip off
        r2 = C.add(args(date=D2, merchant="Ramen", amount=1800, currency="JPY", category="meal", home=15.80, no_trip=True))
        self.assertEqual((r2["home"], r2["trip"]), ("15.80", ""))
        with self.assertRaises(SystemExit):
            C.add(args(merchant="X", category="champagne"))

    def test_add_from_a_json_file(self):
        import contextlib
        import io
        f = C.DATA / "args.json"
        f.write_text(json.dumps({"merchant": "Café \"Rosa\"; $(x)", "amount": 182000, "currency": "VND", "date": D1,
                                 "category": "meal", "description": "Lunch with a client"}), encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(C.main(["add", "--args-json", str(f), "--no-trip"]), 0)
        r = C.load()[0]
        self.assertEqual((r["merchant"], r["home"], r["category"]), ('Café "Rosa"; $(x)', "9.03", "meal"))
        with self.assertRaises(SystemExit):
            C.main(["add", "--currency", "VND"])  # no merchant or amount

    def test_another_home_currency(self):
        C.HOME = "EUR"
        r = C.add(args())
        self.assertEqual(r["home"], "6.30")  # 182,000 VND at 0.0000346 EUR
        self.assertEqual(C.summary(C.load())["currency"], "EUR")
        self.assertEqual(C.rate_text({"currency": "VND", "fx": "0.0000346"}), "1 EUR = 28,902 VND")
        self.assertEqual(C.rate("EUR", D1), 1.0)

    def test_status_summary_and_selection(self):
        a = C.add(args())
        b = C.add(args(date=D2, merchant="Pho 24", amount=450000, category="meal", trip="Hanoi trip"))
        C.set_status([a["id"], b["id"]], "submitted", d(-18))
        C.set_status([a["id"]], "paid", d(-6))
        rows = C.load()
        self.assertEqual(C.select(rows, status="open")[0]["id"], b["id"])
        self.assertEqual(len(C.select(rows, trip="hanoi trip")), 1)
        s = C.summary(rows)
        self.assertEqual(s["by_status"], {"paid": 9.03, "submitted": 22.27})
        self.assertEqual(s["by_trip"]["Hanoi trip"]["open"], 22.27)
        paid = next(r for r in rows if r["id"] == a["id"])
        self.assertEqual((paid["submitted_on"], paid["paid_on"]), (d(-18), d(-6)))
        with self.assertRaises(SystemExit):
            C.set_status(["nope"], "paid", None)

    def test_match_finds_a_reimbursement_for_a_submitted_batch(self):
        a = C.add(args())
        b = C.add(args(date=D2, merchant="Pho 24", amount=450000, category="meal"))
        C.set_status([a["id"], b["id"]], "submitted", d(-18))
        txns = [{"date": d(-25), "amount": "31.31", "desc": "too early"},
                {"date": d(-11), "amount": "31.31", "desc": "EMPLOYER CLAIMS"},
                {"date": d(-10), "amount": "500.00", "desc": "something else"}]
        hits = C.match(C.load(), txns)
        self.assertEqual(len(hits), 1)
        self.assertEqual(hits[0]["txn"]["desc"], "EMPLOYER CLAIMS")
        self.assertEqual(sorted(hits[0]["claims"]), sorted([a["id"], b["id"]]))

    def test_match_without_the_money_bot_is_empty(self):
        real = C.ROOT
        C.ROOT = Path(self.tmp.name) / "no-repo"
        try:
            self.assertEqual(C.reimbursements(), [])
        finally:
            C.ROOT = real

    def test_purge_drops_receipts_60_days_after_paid_only(self):
        dropped = []
        C.drop_receipt = lambda cid: dropped.append(cid)
        old = C.add(args())
        new = C.add(args(date=D2, merchant="Pho 24", amount=450000, category="meal"))
        rows = C.load()
        for r in rows:
            r["receipt"] = "app"
        C.save(rows)
        C.set_status([old["id"]], "paid", d(-70))
        C.set_status([new["id"]], "paid", d(-10))
        gone = C.purge(C.load(), today=TODAY)
        self.assertEqual((gone, dropped), ([old["id"]], [old["id"]]))
        self.assertEqual({r["id"]: r["receipt"] for r in C.load()}, {old["id"]: "", new["id"]: "app"})

    def test_pack_lists_every_line_and_embeds_receipts(self):
        C.add(args())
        rows = C.load()
        rows[0]["receipt"] = "app"
        C.save(rows)
        C.get_receipt = lambda cid: ("image/png", b"\x89PNG fake")
        seen = {}

        def fake_run(cmd, **kw):
            seen["html"] = Path(cmd[2]).read_text()
            Path(cmd[3]).write_bytes(b"%PDF")
            return SimpleNamespace(returncode=0)
        real = C.subprocess.run
        C.subprocess.run = fake_run
        try:
            out = C.pack(C.load(), C.DATA / "claim.pdf", "Expense claim · Hanoi")
        finally:
            C.subprocess.run = real
        self.assertTrue(out.exists())
        self.assertIn("Grab", seen["html"])
        self.assertIn("VND 182,000", seen["html"])
        self.assertIn("data:image/png;base64,", seen["html"])
        self.assertIn("9.03", seen["html"])
        self.assertIn("<th class=r>SGD</th>", seen["html"])
        self.assertFalse((C.DATA / "claim.html").exists())  # the HTML is cleaned up
        with self.assertRaises(SystemExit):
            C.pack([], C.DATA / "empty.pdf", "x")

    def test_export_writes_an_excel_sheet_and_matching_receipts(self):
        import zipfile
        from openpyxl import load_workbook
        a = C.add(args())
        C.add(args(date=D2, merchant="Pho 24 / Hanoi", amount=450000, category="meal"))
        rows = C.load()
        rows[0]["receipt"] = "app"
        C.save(rows)
        C.get_receipt = lambda cid: ("image/jpeg", b"jpeg-bytes") if cid == a["id"] else None
        real_pack = C.pack
        C.pack = lambda rows, out, title: (out.write_bytes(b"%PDF"), out)[1]
        try:
            files = C.export(C.load(), C.DATA / "out", "Hanoi trip")
        finally:
            C.pack = real_pack
        xlsx, pdf, z = files
        self.assertTrue(pdf.name.endswith("_Hanoi trip Claims.pdf"))
        self.assertTrue(xlsx.name.endswith("_Hanoi trip Claims.xlsx"))
        ws = load_workbook(xlsx)["Claims"]
        self.assertEqual([c.value for c in ws[4]][:3], ["No.", "Date", "Merchant"])
        self.assertEqual(ws["H4"].value, "Rate to SGD")
        self.assertEqual(ws["C5"].value, "Grab")
        self.assertEqual(ws["I5"].value, 9.03)
        self.assertEqual(ws["K5"].value, f"01_{D1}_Grab.jpg")
        self.assertEqual(ws["K6"].value, "(none)")
        self.assertEqual(ws["I8"].value, "=SUM(I5:I6)")
        with zipfile.ZipFile(z) as zf:
            self.assertEqual(zf.namelist(), [f"01_{D1}_Grab.jpg"])
            self.assertEqual(zf.read(f"01_{D1}_Grab.jpg"), b"jpeg-bytes")
        with self.assertRaises(SystemExit):
            C.export([], C.DATA / "out", "x")

    def test_a_receipt_photo_is_trimmed_stored_as_jpeg_and_previewed(self):
        import base64
        import io
        import numpy as np
        from PIL import Image
        bg = np.full((800, 600, 3), (90, 65, 45), dtype="uint8")  # a dark table
        bg[150:650, 200:420] = 245  # the slip
        f = C.DATA / "1-receipt.png"
        Image.fromarray(bg).save(f)
        sent = {}
        C.sync.config = lambda: {"url": "https://x"}
        C.sync.token = lambda: "t"
        C.sync.post = lambda cfg, tok, path, obj, max_time=30: sent.update(obj) or {"ok": True}
        cid = f"{D1.replace('-', '')}-abcdef"
        info = C.put_receipt(cid, f)
        self.assertEqual(sent["mime"], "image/jpeg")
        img = Image.open(io.BytesIO(base64.b64decode(sent["data"])))
        self.assertTrue(info["cropped"])
        self.assertLess(img.width, 300)  # the table is cut away: about 220 px of slip plus a margin
        self.assertLess(img.height, 560)
        self.assertTrue(Path(info["preview"]).exists())
        whole = C.put_receipt(cid, f, trim=False)
        self.assertFalse(whole["cropped"])
        by_hand = C.put_receipt(cid, f, crop="0,0,0.5,0.5")
        self.assertEqual(by_hand["how"], "by hand")
        self.assertEqual(C.rate_text({"currency": "VND", "fx": "0.0000496"}), "1 SGD = 20,161 VND")
        self.assertEqual(C.rate_text({"currency": "SGD", "fx": "1"}), "—")

    def test_a_white_receipt_on_a_pale_beige_floor_is_found_by_whiteness(self):
        import numpy as np
        from PIL import Image
        import receipt as R
        bg = np.full((900, 700, 3), (214, 198, 170), dtype="uint8")  # pale beige tiles: bright, but coloured
        bg[600:900, 0:180] = (215, 165, 130)  # a hand
        bg[120:780, 180:520] = (238, 238, 242)  # the receipt
        bg[200:700:40, 220:480] = 60  # lines of print
        box, how = R.find_box(Image.fromarray(bg))
        self.assertEqual(how, "slip")
        for got, want in zip(box, (180 / 700, 120 / 900, 520 / 700, 780 / 900)):
            self.assertAlmostEqual(got, want, delta=0.02)
        self.assertIn("NOT trimmed", C.trim_check("x", {"how": "kept whole", "cropped": False, "preview": "p.jpg"}))

    def test_receipts_need_the_app_and_a_sane_file(self):
        f = C.DATA / "r.txt"
        f.write_text("not an image")
        os.environ.pop("SECRETARY_BOT_TOKEN", None)
        with self.assertRaises(SystemExit):
            C.put_receipt(f"{D1.replace('-', '')}-abcdef", f)


if __name__ == "__main__":
    unittest.main()
